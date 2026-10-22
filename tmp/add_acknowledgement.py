from docx import Document
from docx.shared import Pt

source = 'Crop_Yield_Prediction_DEPSTAR_Format_Completed_Report_Expanded.docx'
output = 'Crop_Yield_Prediction_DEPSTAR_Format_Completed_Report_With_Acknowledgement.docx'
doc = Document(source)
toc_heading = doc.paragraphs[0]

heading = doc.add_paragraph('ACKNOWLEDGEMENT', style='Heading 1')
toc_heading._p.addprevious(heading._p)

ack = doc.add_paragraph(
    'We would like to express our sincere gratitude to our project guide for the guidance, '
    'constructive feedback, and encouragement provided throughout this project. We thank the '
    'faculty members and the Department for their academic support and for providing the resources '
    'needed to complete the work. We also acknowledge the authors and maintainers of the datasets, '
    'software libraries, and technical references used in this project. Finally, we are grateful to '
    'our classmates, friends, and family for their encouragement during the project.'
)
ack.paragraph_format.line_spacing = 1.15
ack.paragraph_format.space_after = Pt(6)
toc_heading._p.addprevious(ack._p)
toc_heading.paragraph_format.page_break_before = True

abstract_entry = next(p for p in doc.paragraphs if p.text.strip().startswith('ABSTRACT'))
toc_entry = doc.add_paragraph('ACKNOWLEDGEMENT')
abstract_entry._p.addprevious(toc_entry._p)  # before ABSTRACT in the contents list

doc.save(output)
print('Saved:', output)
