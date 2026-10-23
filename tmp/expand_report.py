from docx import Document
from docx.shared import Pt
from docx.shared import Inches

SOURCE = 'Crop_Yield_Prediction_DEPSTAR_Format_Completed_Report_Updated.docx'
PATH = 'Crop_Yield_Prediction_DEPSTAR_Format_Completed_Report_Expanded.docx'
doc = Document(SOURCE)

def add_after_prefix(prefix, text):
    anchor = next((p for p in doc.paragraphs if p.text.strip().startswith(prefix)), None)
    if anchor is None:
        raise RuntimeError(f'Anchor not found: {prefix}')
    new_p = doc.add_paragraph()
    new_p.style = doc.styles['Normal']
    new_p.paragraph_format.line_spacing = 1.15
    new_p.paragraph_format.space_after = Pt(6)
    new_p.add_run(text)
    anchor._p.addnext(new_p._p)

add_after_prefix('The completed implementation covers',
    'From a user perspective, the prototype has two complementary roles. The Home page helps a reader inspect patterns already present in the project records, while the Crop Prediction page turns a selected set of scenario inputs into an estimated yield. This distinction matters because a historical summary describes the stored dataset and its filters; it does not describe a newly observed field or a live crop condition. The intended audience is a student, evaluator or analyst exploring the project and its model behavior.')

add_after_prefix('The system requires local data files',
    'The non-functional requirements are modest and appropriate to an academic prototype. The interface should remain understandable on a standard desktop browser, present units beside numeric outputs, and load saved artifacts without requiring model training during each interaction. Results should be reproducible from the checked-in data and model files when compatible dependencies are installed. The project is not sized or evaluated as a production service: uptime, concurrent-user capacity, security hardening and operational support have not been established.')

add_after_prefix('The source crop_yield.csv contains',
    'The state-selection helper shown in Figure 3.1 is a preprocessing utility that extracts a Rice and Maize subset for a reference lookup workflow. Its narrower crop list is not the crop list of the completed application or full modeling benchmark. The broader application and tabular results support Rice, Maize, Moong (Green Gram) and Urad. Keeping that distinction explicit prevents a data-preparation screenshot from being mistaken for the project-wide supported-crop scope.')

add_after_prefix('The documented modeling workflow removes',
    'Before modeling, identifiers and field names must be normalized so joins do not silently fail because of spacing or spelling differences. The working tables are checked for missing target values, duplicate keys and implausible numeric entries. Yield is treated as the response variable, and production is excluded from predictor features because yield is calculated from production and area. A feature that directly reconstructs the target would inflate evaluation scores without representing a valid forecast input.')

add_after_prefix('Python is used for data preparation',
    'The implementation is organized around reusable artifacts rather than a training action inside the interface. Preprocessing and experiment scripts prepare features and partitions; evaluation scripts save predictions and metric summaries; and app.py loads the required data and estimators for interactive use. This separation shortens dashboard response time and makes the reported experiment outputs traceable to repository files. Reproduction still depends on using a compatible Python environment and retaining the referenced datasets, image files and serialized models.')

add_after_prefix('The pipeline begins with',
    'At inference time, the tabular path assembles a row from the selected state, crop, season, year and numeric scenario values, applies the saved preprocessing/model path, and returns a yield estimate. The model output is an estimate in tonnes per hectare, conditioned on the trained model and available feature representation. It should be interpreted alongside the input assumptions, rather than as a guaranteed harvest or a prescriptive action. The project does not provide confidence intervals or an automated agronomic decision.')

add_after_prefix('The current Streamlit sidebar exposes',
    'The dashboard’s Home view is arranged for quick exploration: filters sit above four summary cards, followed by a compact data-context strip. Below it, the application places the state distribution map above the top-states chart in the left column, and the selected-metric trend above crop-area distribution in the right column. The metric selector changes which stored measure is emphasized, while year, state and crop controls narrow the displayed records. Figure 5.2 follows this arrangement; chart values in the live application are calculated from the active filters.')

add_after_prefix('The modeling table covers Rice',
    'The four-crop comparison shows that both boosting models substantially outperform the dummy-mean reference on this split. The validation-selected HistGradientBoosting model has a slightly lower test R² than Gradient Boosting, but a lower test MAE; model selection remains tied to validation performance to avoid choosing a model after inspecting the held-out test score. The difference between the two boosting test scores is small, so the report does not claim a universal winner beyond the selected validation protocol.')

add_after_prefix('The modeling table covers Rice',
    'Crop-wise R² is uneven: the highest value among the four crops is for Maize (0.6513), followed by Rice (0.5332) and Moong (0.3716), while Urad is much lower (0.0374). Each crop has between 44 and 63 test records. This variation is a reason to report per-crop behavior alongside the aggregate score: a strong pooled result can hide weaker performance for a particular target crop. The values are specific to the available 2019–2020 test records.')

add_after_prefix('The completed artifacts show that',
    'The available evidence supports a focused set of functional checks: the configured target list contains the four intended crops; dashboard filters are populated from project records; the prediction form collects the documented scenario fields; and saved estimators and metric files are present for the experiments described. Data integrity review also confirms that the target is kept separate from model predictors and that image experiments use aligned RGB/NDVI pairs. These checks document observed project behavior; they are not a claim of a formal automated test suite or production acceptance testing.')

add_after_prefix('The main engineering challenge was',
    'A further limitation is that state-level aggregation cannot resolve within-state differences in irrigation, soil management, cultivar, planting date or local weather. Likewise, a stored state-level image sample may not correspond to the same plot, season or exact observation date as a yield record. The report therefore describes the imagery as aligned experiment input where the repository manifest defines a pair, and as stored preview material in the dashboard. Neither use supports field-level diagnosis.')

add_after_prefix('Future work can build a larger',
    'A useful next evaluation would define one shared complete-case cohort with matching tabular, satellite-summary and image inputs, then preserve the same chronological test records for each model arm. Reporting sample counts per crop, state and year would make comparisons easier to audit. Repeated temporal or spatial holdouts could show whether a measured gain persists beyond one partition. Until those comparisons are run and documented, the existing fusion increase should remain described as a small observed difference on its saved paired evaluation.')

# Put the app.py-faithful dashboard layout illustration into the previously empty
# figure slot. It is explicitly labelled as a recreation, not a captured screenshot.
caption = next(p for p in doc.paragraphs if p.text.strip().startswith('Fig. 5.2'))
previous = caption._p.getprevious()
if previous is None:
    raise RuntimeError('Dashboard figure insertion point is missing')
from docx.text.paragraph import Paragraph
slot = Paragraph(previous, caption._parent)
slot.clear()
slot.alignment = 1
slot.add_run().add_picture('tmp/report_assets/dashboard_illustration.png', width=Inches(6.25))
caption.text = 'Fig. 5.2   Dashboard Layout Recreated from app.py'
note = next((p for p in doc.paragraphs if p.text.strip().startswith('Recreated from app.py.')), None)
if note is not None:
    note.text = ('Layout recreation based on the current app.py Home page. This is not a captured browser screenshot; '
                 'the app calculates chart values from the active filters, and the illustration omits those data values.')

doc.save(PATH)
print('Expanded report saved:', PATH)
